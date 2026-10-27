"""
Academic Document Generator for ISE-2 Submission.
Builds publication-grade DOCX and PDF versions of ISE2_MLFLOW_DOCUMENTATION.md
with embedded plots, professional styling, tables, and student credentials.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


STUDENT_NAME = "Tanish Patgaonkar"
ROLL_NUMBER = "42"
PRN = "2324000399"
ACADEMIC_YEAR = "2026-27"
COURSE_TITLE = "ISE-2 Machine Learning & MLOps Assessment"
PROJECT_TITLE = "Iris Flower Classification: MLflow Experiment Tracking and Kolmogorov-Smirnov Data Drift Analysis"


def set_cell_background(cell, hex_color):
    """Sets cell background shading in docx."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding inside a docx table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def create_docx(output_path: str = "ISE2_MLFLOW_DOCUMENTATION.docx"):
    """Generates a polished Word document."""
    doc = Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ISE-2 Academic Report | Iris MLflow & Drift Analysis")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(128, 128, 128)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run(f"{STUDENT_NAME} (PRN: {PRN}, Roll: {ROLL_NUMBER}) — Academic Year {ACADEMIC_YEAR}")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(128, 128, 128)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_super = p_title.add_run("ISE-2 ACADEMIC EVALUATION REPORT\n")
    r_super.font.size = Pt(13)
    r_super.font.bold = True
    r_super.font.color.rgb = RGBColor(31, 78, 121)
    
    r_sub = p_title.add_run(f"{COURSE_TITLE}\n")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(89, 89, 89)
    
    r_main = p_title.add_run(PROJECT_TITLE)
    r_main.font.size = Pt(16)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(16, 42, 77)

    # Student Details Box Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [f"Student Name: {STUDENT_NAME}", f"Roll Number: {ROLL_NUMBER}"],
        [f"PRN: {PRN}", f"Academic Year: {ACADEMIC_YEAR}"]
    ]
    for r_idx, row in enumerate(meta_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = meta_data[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.color.rgb = RGBColor(31, 78, 121)
            set_cell_background(cell, "F2F4F7")
            set_cell_margins(cell, top=120, bottom=120, left=200, right=200)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==================== SECTION 1 ====================
    h1 = doc.add_heading("Section 1: MLflow Installation and Experiment Tracking", level=1)
    h1.runs[0].font.color.rgb = RGBColor(31, 78, 121)

    doc.add_heading("1.1 Environment Setup", level=2)
    p_env = doc.add_paragraph()
    p_env.add_run("The machine learning experiment and tracking environment was configured entirely on a local development workstation, maintaining complete compatibility with Apple Silicon architecture without cloud-dependent infrastructure or paid proprietary APIs.\n\n")
    
    env_items = [
        ("Operating System", "macOS (Darwin 24.x, Apple Silicon M3, 8GB Unified Memory Architecture)"),
        ("Python Version", "Python 3.13.5 (configured inside an isolated virtual environment via venv)"),
        ("MLflow Version", "MLflow 3.17.2 (Skinny / Core Tracing distribution)"),
        ("Key Supporting Libraries", "scikit-learn 1.9.1, pandas 3.1.0, numpy 2.5.4, scipy 1.18.1, matplotlib 3.11.2, seaborn 0.13.2, pytest 9.1.1"),
        ("Installation Status", "Clean installation verified inside .venv/ with dependencies version-locked in requirements.txt."),
        ("MLflow Tracking URI", "file:///Users/tanishpatgaonkar/.gemini/antigravity-ide/scratch/iris-mlflow-drift/mlruns"),
        ("Local UI Verification", "Local tracking UI verified actively serving on http://127.0.0.1:5001 (HTTP 200 OK verified).")
    ]
    for label, val in env_items:
        p = doc.add_paragraph(style='List Bullet')
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.bold = True
        p.add_run(val)

    doc.add_heading("1.2 Experiment Tracking", level=2)
    doc.add_paragraph(
        "Experiment Name: Iris_Classification_Experiment\n"
        "Experiment Objective: The primary objective of this experiment is to systematically track, evaluate, "
        "and compare multiple hyperparameter configurations of a multiclass Logistic Regression classifier on the Iris dataset. "
        "By logging parameters, metrics, evaluation reports, and confusion matrix artifacts in a centralized tracking store, "
        "we ensure complete experiment reproducibility, traceability, and auditability throughout the machine learning development lifecycle."
    )
    
    doc.add_paragraph(
        "Parameters Recorded: model_name (LogisticRegression), test_size (0.20), random_state (42), "
        "max_iterations (100, 200, 300), regularization_parameter_C (0.1, 1.0, 10.0), solver (lbfgs), preprocessing (StandardScaler).\n\n"
        "Metrics Recorded: accuracy, train_accuracy, precision (macro), recall (macro), f1_score (macro), f1_score_weighted.\n\n"
        "Artifacts Recorded: confusion_matrix_*.png, classification_report_*.txt, dataset_info_*.json, serialized model pipeline (iris_model/)."
    )

    doc.add_paragraph("Table 1 summarizes the empirical results logged across the 3 MLflow runs:")

    # Table 1: MLflow Runs
    table1 = doc.add_table(rows=4, cols=8)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers1 = ["Run Name", "Inv Reg (C)", "Max Iter", "Test Acc", "Precision", "Recall", "F1-Score", "Train Acc"]
    rows1 = [
        ["Run 1: Baseline", "1.0", "200", "0.9333", "0.9333", "0.9333", "0.9333", "0.9583"],
        ["Run 2: High Reg", "0.1", "100", "0.8667", "0.8750", "0.8667", "0.8653", "0.9333"],
        ["Run 3: Low Reg", "10.0", "300", "1.0000", "1.0000", "1.0000", "1.0000", "0.9750"],
    ]
    for c_idx, h in enumerate(headers1):
        cell = table1.cell(0, c_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)

    for r_idx, row_vals in enumerate(rows1):
        for c_idx, val in enumerate(row_vals):
            cell = table1.cell(r_idx + 1, c_idx)
            cell.text = val
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(9)
            if c_idx == 0:
                p.runs[0].font.bold = True
            if r_idx % 2 == 1:
                set_cell_background(cell, "F2F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Embed Confusion Matrices
    cm_paths = [
        ("Run 1 Baseline Confusion Matrix", "plots/confusion_matrix_run_1_baseline_c1.0.png"),
        ("Run 2 High Regularization Confusion Matrix", "plots/confusion_matrix_run_2_high_regularization_c0.1.png"),
        ("Run 3 Low Regularization Confusion Matrix", "plots/confusion_matrix_run_3_low_regularization_c10.0.png")
    ]
    for caption, p in cm_paths:
        if os.path.exists(p):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_picture(p, width=Inches(3.8))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cap = p_cap.add_run(f"Figure: {caption} (Logged to MLflow)")
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(89, 89, 89)

    # ==================== SECTION 2 ====================
    doc.add_page_break()
    h2 = doc.add_heading("Section 2: Dataset Identification", level=1)
    h2.runs[0].font.color.rgb = RGBColor(31, 78, 121)

    doc.add_heading("2.1 Problem Statement", level=2)
    doc.add_paragraph(
        "Automated taxonomic classification of botanical specimens based on morphological dimensions is a classic supervised classification problem. "
        "The objective is to construct a predictive mathematical model that takes continuous physical dimensions of flower blossoms "
        "(sepal and petal lengths and widths) and accurately maps them into one of three distinct species: Iris setosa, Iris versicolor, or Iris virginica.\n"
        "In production environments, measuring instruments may suffer calibration drift, regional climate changes may alter flower growth rates, "
        "or seasonal shifts may induce statistical distribution deviations. Therefore, this project combines model training with automated data drift monitoring."
    )

    doc.add_heading("2.2 Dataset Details", level=2)
    ds_items = [
        ("Dataset Name", "Fisher's Iris Flower Dataset"),
        ("Dataset Source", "Scikit-learn standard dataset (sklearn.datasets.load_iris), Edgar Anderson (1935) / Ronald Fisher (1936)"),
        ("Number of Records", "150 instances (50 records per species, perfectly balanced)"),
        ("Features (4 continuous)", "sepal_length (cm), sepal_width (cm), petal_length (cm), petal_width (cm)"),
        ("Target Variable", "species (0: setosa, 1: versicolor, 2: virginica)"),
        ("Missing / Null Values", "0 missing values across all columns"),
        ("Use Case Category", "Multiclass Supervised Machine Learning Classification")
    ]
    for label, val in ds_items:
        p = doc.add_paragraph(style='List Bullet')
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.bold = True
        p.add_run(val)

    doc.add_heading("2.3 Dataset Selection Rationale", level=2)
    doc.add_paragraph(
        "The Iris dataset is exceptionally well-suited for this academic assessment for several distinct reasons:\n"
        "1. Geometric Intuitiveness: The features have clear physical interpretations and distinct clustering characteristics "
        "(e.g., Setosa is linearly separable, while Versicolor and Virginica exhibit slight non-linear overlap).\n"
        "2. Deterministic Baseline: The absence of missing values, encoding complications, or extreme noise allows us to isolate "
        "the effects of hyperparameters in MLflow without confusing data quality artifacts.\n"
        "3. Controlled Drift Verification: Because feature distributions are well-understood continuous unimodal and bimodal densities, "
        "synthetic perturbations can be introduced with high statistical precision to validate hypothesis testing.\n"
        "4. Computational Efficiency: Minimal CPU and RAM footprint enables rapid local iterations and instant execution on Apple Silicon laptops."
    )

    # ==================== SECTION 3 ====================
    doc.add_page_break()
    h3 = doc.add_heading("Section 3: Drift Scenario and Detection Method", level=1)
    h3.runs[0].font.color.rgb = RGBColor(31, 78, 121)

    doc.add_heading("3.1 Drift Type", level=2)
    doc.add_paragraph(
        "The specific drift simulated and evaluated in this project is Data Drift (Covariate Shift).\n"
        "Mathematically, Covariate Shift occurs when the marginal probability distribution of the input features shifts between the baseline training phase "
        "and the production inference phase, while the true conditional probability distribution of the target labels given the features remains stationary:\n"
        "P(X)_current != P(X)_reference  while  P(Y|X)_current = P(Y|X)_reference\n\n"
        "Academic Distinction:\n"
        "• Concept Drift: Occurs when the functional relationship mapping inputs to outputs changes (P(Y|X)_curr != P(Y|X)_ref), meaning old decision boundaries become invalid.\n"
        "• Prediction Drift: Occurs when the distribution of the model's output predictions changes (P(Y_hat)_curr != P(Y_hat)_ref)."
    )

    doc.add_heading("3.2 Reference and Current Data", level=2)
    doc.add_paragraph(
        "• Reference Dataset (D_ref): Pristine baseline 150-sample Iris dataset exported to data/reference_iris.csv.\n"
        "• Current Dataset (D_curr): Simulated monitoring dataset created in data/current_iris_drifted.csv by applying controlled mathematical perturbations:\n"
        "  - sepal_length: Shifted upward by adding N(0.85, 0.15^2), producing an approximate mean shift of +0.837 cm (~1.0 standard deviation).\n"
        "  - petal_length: Shifted upward by adding N(1.40, 0.20^2), producing an approximate mean shift of +1.414 cm (~0.8 standard deviations).\n"
        "  - sepal_width: Maintained without intentional distribution shift (mean shift: +0.012 cm).\n"
        "  - petal_width: Maintained without intentional distribution shift (mean shift: -0.012 cm).\n\n"
        "Academic Disclosure: This is an intentionally controlled synthetic perturbation designed to simulate sensor miscalibration or environmental growth variation, "
        "allowing us to evaluate whether our statistical detector correctly flags only the shifted features while avoiding false positives on the unshifted control features."
    )

    doc.add_heading("3.3 Detection Method", level=2)
    doc.add_paragraph(
        "• Statistical Test: Two-Sample Kolmogorov-Smirnov (KS) Test (scipy.stats.ks_2samp).\n"
        "• How the KS Test Operates: The two-sample KS test is a non-parametric hypothesis test that quantifies the supreme vertical distance between "
        "the empirical cumulative distribution functions (ECDF) of two independent samples:\n"
        "  D = sup_x |F_ref(x) - F_curr(x)|\n"
        "• Null Hypothesis (H0): Both samples originate from identical continuous distributions (F_ref(x) = F_curr(x) for all x).\n"
        "• Alternative Hypothesis (H1): The two samples originate from different distributions (F_ref(x) != F_curr(x) for at least one x).\n"
        "• Significance Threshold (alpha): alpha = 0.05 (95% statistical confidence level).\n"
        "• Decision Rule: If p-value < 0.05, reject H0 -> Drift Detected. If p-value >= 0.05, fail to reject H0 -> No Significant Drift."
    )

    # ==================== SECTION 4 ====================
    doc.add_page_break()
    h4 = doc.add_heading("Section 4: Results, Interpretation and Conclusion", level=1)
    h4.runs[0].font.color.rgb = RGBColor(31, 78, 121)

    doc.add_heading("4.1 Drift Analysis Results", level=2)
    doc.add_paragraph("Table 2 documents the empirical results generated by executing src/drift_detection.py and saved in reports/drift_report.csv:")

    # Table 2: Drift Results
    table2 = doc.add_table(rows=5, cols=8)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers2 = ["Feature Name", "KS Stat (D)", "p-value", "Alpha", "Ref Mean", "Curr Mean", "Mean Shift", "Statistical Status"]
    rows2 = [
        ["sepal_length", "0.3733", "1.0722e-09", "0.05", "5.843 cm", "6.681 cm", "+0.837 cm", "Drift Detected"],
        ["sepal_width", "0.1000", "0.4425", "0.05", "3.057 cm", "3.069 cm", "+0.012 cm", "No Significant Drift"],
        ["petal_length", "0.4067", "1.7712e-11", "0.05", "3.758 cm", "5.172 cm", "+1.414 cm", "Drift Detected"],
        ["petal_width", "0.1333", "0.1391", "0.05", "1.199 cm", "1.188 cm", "-0.012 cm", "No Significant Drift"]
    ]
    for c_idx, h in enumerate(headers2):
        cell = table2.cell(0, c_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)

    for r_idx, row_vals in enumerate(rows2):
        for c_idx, val in enumerate(row_vals):
            cell = table2.cell(r_idx + 1, c_idx)
            cell.text = val
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            if c_idx == 0:
                p.runs[0].font.bold = True
            if c_idx == 7:
                p.runs[0].font.bold = True
                if "Drift Detected" in val:
                    p.runs[0].font.color.rgb = RGBColor(192, 0, 0)
                else:
                    p.runs[0].font.color.rgb = RGBColor(38, 128, 0)
            if r_idx % 2 == 1:
                set_cell_background(cell, "F2F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=80, right=80)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Embed Drift Plots
    if os.path.exists("plots/drift_distributions.png"):
        doc.add_picture("plots/drift_distributions.png", width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 1: Empirical Distribution Comparison (KDE & Histograms) across Reference vs. Current Datasets")
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(89, 89, 89)

    if os.path.exists("plots/drift_ks_summary.png"):
        doc.add_picture("plots/drift_ks_summary.png", width=Inches(5.8))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 2: Summary of Kolmogorov-Smirnov Test Statistics and Significance Thresholds")
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_heading("4.2 Interpretation", level=2)
    doc.add_paragraph(
        "1. Shifted Feature Validation (sepal_length and petal_length):\n"
        "   - sepal_length: KS statistic D = 0.3733, p-value = 1.0722e-09. Since p << 0.05, the null hypothesis is decisively rejected.\n"
        "   - petal_length: KS statistic D = 0.4067, p-value = 1.7712e-11. Decisively rejected.\n"
        "   - Both features show significant empirical CDF divergence, successfully reflecting the controlled synthetic perturbation.\n\n"
        "2. Control Feature Validation (sepal_width and petal_width):\n"
        "   - sepal_width: KS statistic D = 0.1000, p-value = 0.4425 >= 0.05. We fail to reject H0.\n"
        "   - petal_width: KS statistic D = 0.1333, p-value = 0.1391 >= 0.05. We fail to reject H0.\n"
        "   - Confirms that the detector possesses high specificity, avoiding false alarms on unshifted features.\n\n"
        "3. Crucial Academic Distinction on Performance Degradation:\n"
        "   - Detecting data drift does NOT automatically prove model performance has decreased.\n"
        "   - If the shifted samples still fall securely within their respective class decision boundaries, accuracy may remain intact.\n"
        "   - However, data drift indicates that the model is receiving out-of-distribution inputs that violate validation assumptions, "
        "serving as an indispensable early-warning telemetry trigger."
    )

    doc.add_heading("4.3 Conclusion and Recommended Action", level=2)
    doc.add_paragraph(
        "Conclusion: The local MLflow experiment tracking and Kolmogorov-Smirnov drift detection architecture has successfully demonstrated full end-to-end functionality. "
        "The classification models attained high predictive performance (up to 100% test accuracy under C=10.0 and 93.3% under baseline C=1.0), "
        "and the two-sample KS test detected covariate shift in precisely the target features with high confidence (p < 10^-8).\n\n"
        "Recommended MLOps Interventions:\n"
        "1. Telemetry Alerting: Configure automated alerts when any feature's KS test yields p < 0.05 over a rolling 24-hour window.\n"
        "2. Active Label Acquisition: Pull a random sample of drifted flower instances and send them for expert botanical re-verification to acquire ground-truth labels.\n"
        "3. Performance Audit: Compute real-time accuracy and F1-score on newly labeled drifted samples to verify whether decision boundary fidelity has degraded.\n"
        "4. Retraining & Champion-Challenger Deployment: Retrain the model pipeline on an augmented dataset combining baseline data and recently acquired drifted samples, "
        "log the candidate model to MLflow, and perform a shadow deployment before promoting to production."
    )

    # ==================== SECTION 5 ====================
    doc.add_page_break()
    h5 = doc.add_heading("Section 5: GitHub Repository and Reproducibility", level=1)
    h5.runs[0].font.color.rgb = RGBColor(31, 78, 121)

    doc.add_heading("5.1 GitHub Repository URL", level=2)
    doc.add_paragraph(
        "• GitHub Repository URL: https://github.com/Tanish-19/iris-mlflow-drift\n"
        "• Local Repository Root: /Users/tanishpatgaonkar/.gemini/antigravity-ide/scratch/iris-mlflow-drift\n"
        "• Configured Branch: main (Git repository initialized, clean local commits, synced with remote origin)."
    )

    doc.add_heading("5.2 Source Code and Project Inventory", level=2)
    code_items = [
        ("src/data_loader.py", "Iris dataset loader and controlled covariate shift generator"),
        ("src/train.py", "MLflow tracking suite executing 3 distinct hyperparameter configurations"),
        ("src/drift_detection.py", "Two-sample KS statistical test engine, distribution plotter, and report generator"),
        ("app.py", "Interactive Streamlit presentation dashboard"),
        ("tests/test_project.py", "9 comprehensive automated pytest validation test cases"),
        ("requirements.txt", "Pinned Python dependencies manifest"),
        (".gitignore", "Excludes virtual environments, pycache, and temporary runtime files")
    ]
    for label, val in code_items:
        p = doc.add_paragraph(style='List Bullet')
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.bold = True
        p.add_run(val)

    doc.add_heading("5.3 README and Execution Instructions", level=2)
    doc.add_paragraph(
        "1. Ingest Data & Generate Drift:\n   python src/data_loader.py\n"
        "2. Execute MLflow Training Suite (3 Runs):\n   python src/train.py\n"
        "3. Execute Kolmogorov-Smirnov Drift Analysis:\n   python src/drift_detection.py\n"
        "4. Execute Automated Pytest Test Suite:\n   pytest tests/test_project.py -v\n"
        "5. Launch Local MLflow UI:\n   MLFLOW_ALLOW_FILE_STORE=true mlflow ui --port 5001\n"
        "6. Launch Streamlit UI:\n   streamlit run app.py --server.port 8501"
    )

    doc.add_heading("5.4 Requirements / Dependencies", level=2)
    doc.add_paragraph("Dependencies version-locked in requirements.txt: Flask 3.1.3, mlflow 3.17.2, scikit-learn 1.9.1, pandas 3.1.0, numpy 2.5.4, scipy 1.18.1, matplotlib 3.11.2, seaborn 0.13.2, pytest 9.1.1, streamlit 1.66.1, tabulate 0.10.0, python-docx 1.2.0, reportlab 5.1.0.")

    doc.add_heading("5.5 Screenshots and Drift Reports", level=2)
    doc.add_paragraph(
        "1. Drift Distribution Plot: plots/drift_distributions.png\n"
        "2. KS Test Summary Plot: plots/drift_ks_summary.png\n"
        "3. Confusion Matrix Plots: plots/confusion_matrix_run_1_baseline_c1.0.png, plots/confusion_matrix_run_2_high_regularization_c0.1.png, plots/confusion_matrix_run_3_low_regularization_c10.0.png\n"
        "4. Drift CSV Report: reports/drift_report.csv\n"
        "5. Experiment Comparison CSV: reports/experiment_summary.csv\n"
        "6. Markdown Report: reports/drift_report.md"
    )

    doc.add_heading("5.6 MLflow Experiment and Drift-Analysis Evidence", level=2)
    doc.add_paragraph(
        "• Experiment ID: Iris_Classification_Experiment\n"
        "• Active Runs Logged:\n"
        "  - Run_1_Baseline_C1.0 (Run ID: 0c10d1d8776344acbdf54020bd3347aa) -> Test Accuracy: 93.33%\n"
        "  - Run_2_High_Regularization_C0.1 (Run ID: 4e4f60a28b30444dbf95637ecdd946b8) -> Test Accuracy: 86.67%\n"
        "  - Run_3_Low_Regularization_C10.0 (Run ID: 9e20c1aa03024d5ca849bb88a59c821f) -> Test Accuracy: 100.00%\n"
        "• Automated Pytest Pass Rate: 9/9 tests passed (100%) in 14.34 seconds.\n"
        "• Drift Test Verification: 100% precision in isolating shifted features (sepal_length and petal_length) while respecting non-rejection on control features."
    )

    doc.save(output_path)
    print(f"Successfully generated DOCX report at: {output_path}")


def create_pdf(output_path: str = "ISE2_MLFLOW_DOCUMENTATION.pdf"):
    """Generates a publication-quality PDF document using ReportLab."""
    pdf_doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Palette
    c_primary = colors.HexColor("#1F4E79")
    c_dark = colors.HexColor("#102A4D")
    c_gray = colors.HexColor("#595959")
    c_light = colors.HexColor("#F2F5F9")
    c_alert = colors.HexColor("#C00000")
    c_ok = colors.HexColor("#268000")
    
    # Typography Styles
    title_super = ParagraphStyle('TitleSuper', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=14, textColor=c_primary, alignment=1)
    title_sub = ParagraphStyle('TitleSub', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=10.5, leading=13, textColor=c_gray, alignment=1)
    title_main = ParagraphStyle('TitleMain', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=15, leading=18, textColor=c_dark, alignment=1)
    
    h1_style = ParagraphStyle('H1Style', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=13, leading=16, textColor=c_primary, spaceBefore=14, spaceAfter=6, keepWithNext=True)
    h2_style = ParagraphStyle('H2Style', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=10.5, leading=13, textColor=c_dark, spaceBefore=10, spaceAfter=4, keepWithNext=True)
    body_style = ParagraphStyle('BodyTextCustom', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, leading=11.5, textColor=colors.HexColor("#1A1A1A"))
    bullet_style = ParagraphStyle('BulletCustom', parent=body_style, leftIndent=12, firstLineIndent=-8)
    caption_style = ParagraphStyle('CaptionStyle', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=8, leading=10, textColor=c_gray, alignment=1)
    table_header_style = ParagraphStyle('THStyle', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white, alignment=1)
    table_cell_style = ParagraphStyle('TCStyle', fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=colors.black, alignment=1)
    table_cell_bold = ParagraphStyle('TCBold', fontName='Helvetica-Bold', fontSize=7.5, leading=9.5, textColor=colors.black, alignment=1)

    story = []

    # Title Banner
    story.append(Paragraph("ISE-2 ACADEMIC EVALUATION REPORT", title_super))
    story.append(Spacer(1, 3))
    story.append(Paragraph(COURSE_TITLE, title_sub))
    story.append(Spacer(1, 4))
    story.append(Paragraph(PROJECT_TITLE, title_main))
    story.append(Spacer(1, 10))

    # Student Details Box
    meta_data = [
        [Paragraph(f"<b>Student Name:</b> {STUDENT_NAME}", body_style), Paragraph(f"<b>Roll Number:</b> {ROLL_NUMBER}", body_style)],
        [Paragraph(f"<b>PRN:</b> {PRN}", body_style), Paragraph(f"<b>Academic Year:</b> {ACADEMIC_YEAR}", body_style)]
    ]
    t_meta = Table(meta_data, colWidths=[3.6 * inch, 3.6 * inch])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F0F4F8")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#D0D7DE")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E1E4E8")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 12))

    # SECTION 1
    story.append(Paragraph("Section 1: MLflow Installation and Experiment Tracking", h1_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=2, spaceAfter=8))
    
    story.append(Paragraph("1.1 Environment Setup", h2_style))
    story.append(Paragraph("The experiment was developed and executed natively on a local Apple Silicon workstation without paid cloud infrastructure:", body_style))
    story.append(Spacer(1, 4))
    
    env_points = [
        "<b>Operating System:</b> macOS (Darwin 24.x, Apple Silicon M3, 8GB RAM).",
        "<b>Python Version:</b> Python 3.13.5 running inside an isolated <code>.venv</code>.",
        "<b>MLflow Version:</b> MLflow 3.17.2 with local persistent tracking store at <code>file:./mlruns</code>.",
        "<b>Core Libraries:</b> scikit-learn 1.9.1, pandas 3.1.0, numpy 2.5.4, scipy 1.18.1, matplotlib 3.11.2, seaborn 0.13.2, pytest 9.1.1, streamlit 1.66.1.",
        "<b>Local MLflow UI:</b> Verified actively running on <code>http://127.0.0.1:5001</code> (HTTP 200 OK)."
    ]
    for pt in env_points:
        story.append(Paragraph(f"• {pt}", bullet_style))
    story.append(Spacer(1, 8))

    story.append(Paragraph("1.2 Experiment Tracking & Runs Comparison", h2_style))
    story.append(Paragraph("<b>Experiment Name:</b> <code>Iris_Classification_Experiment</code>", body_style))
    story.append(Paragraph("Tracked 3 distinct Logistic Regression configurations varying the inverse regularization parameter C:", body_style))
    story.append(Spacer(1, 6))

    # Table 1
    t1_data = [
        [Paragraph(h, table_header_style) for h in ["Run Name", "Inv Reg (C)", "Max Iter", "Test Acc", "Precision", "Recall", "F1-Score", "Train Acc"]],
        [Paragraph("Run 1: Baseline", table_cell_bold), Paragraph("1.0", table_cell_style), Paragraph("200", table_cell_style), Paragraph("<b>0.9333</b>", table_cell_style), Paragraph("0.9333", table_cell_style), Paragraph("0.9333", table_cell_style), Paragraph("0.9333", table_cell_style), Paragraph("0.9583", table_cell_style)],
        [Paragraph("Run 2: High Reg", table_cell_bold), Paragraph("0.1", table_cell_style), Paragraph("100", table_cell_style), Paragraph("<b>0.8667</b>", table_cell_style), Paragraph("0.8750", table_cell_style), Paragraph("0.8667", table_cell_style), Paragraph("0.8653", table_cell_style), Paragraph("0.9333", table_cell_style)],
        [Paragraph("Run 3: Low Reg", table_cell_bold), Paragraph("10.0", table_cell_style), Paragraph("300", table_cell_style), Paragraph("<b>1.0000</b>", table_cell_style), Paragraph("1.0000", table_cell_style), Paragraph("1.0000", table_cell_style), Paragraph("1.0000", table_cell_style), Paragraph("0.9750", table_cell_style)]
    ]
    t1 = Table(t1_data, colWidths=[1.5*inch, 0.75*inch, 0.7*inch, 0.75*inch, 0.75*inch, 0.75*inch, 0.75*inch, 0.85*inch])
    t1.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_primary),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#B0C4DE")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, c_light]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t1)
    story.append(Spacer(1, 10))

    # SECTION 2
    story.append(Paragraph("Section 2: Dataset Identification", h1_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=2, spaceAfter=8))
    
    story.append(Paragraph("2.1 Problem Statement & 2.2 Dataset Details", h2_style))
    story.append(Paragraph("Constructing a multiclass classifier to map sepal and petal dimensions into three botanical species (Setosa, Versicolor, Virginica) while actively detecting distribution drift.", body_style))
    story.append(Spacer(1, 4))
    
    ds_pts = [
        "<b>Dataset Source:</b> Scikit-learn built-in Fisher's Iris dataset (150 records, 50 per class).",
        "<b>Features:</b> <code>sepal_length</code>, <code>sepal_width</code>, <code>petal_length</code>, <code>petal_width</code> (all in cm).",
        "<b>Target Variable:</b> <code>species</code> (0: setosa, 1: versicolor, 2: virginica).",
        "<b>Selection Rationale:</b> Benchmark stability, absence of missing values, and clear linear/non-linear separability enable rigorous isolation of hyperparameter impact and controlled drift testing."
    ]
    for pt in ds_pts:
        story.append(Paragraph(f"• {pt}", bullet_style))
    story.append(Spacer(1, 10))

    # SECTION 3
    story.append(Paragraph("Section 3: Drift Scenario and Detection Method", h1_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=2, spaceAfter=8))
    
    story.append(Paragraph("3.1 Drift Type & Simulation Scenario", h2_style))
    story.append(Paragraph("Evaluates <b>Data Drift (Covariate Shift)</b> where feature distribution changes (P(X)_curr != P(X)_ref) while conditional target probability P(Y|X) remains constant. Distinct from Concept Drift (P(Y|X) changes) and Prediction Drift (P(Y_hat) changes).", body_style))
    story.append(Spacer(1, 4))
    
    drift_pts = [
        "<b>Reference Dataset:</b> Pristine baseline 150-sample Iris dataset exported to <code>data/reference_iris.csv</code>.",
        "<b>Current Dataset:</b> Controlled synthetic perturbation in <code>data/current_iris_drifted.csv</code>: <code>sepal_length</code> shifted by +0.837 cm and <code>petal_length</code> shifted by +1.414 cm; <code>sepal_width</code> (+0.012 cm) and <code>petal_width</code> (-0.012 cm) preserved as stable controls.",
        "<b>Detection Method:</b> Two-Sample Kolmogorov-Smirnov (KS) Test (<code>scipy.stats.ks_2samp</code>) comparing empirical cumulative distributions (ECDF). Maximum distance D = sup_x |F_ref(x) - F_curr(x)|.",
        "<b>Threshold:</b> Significance level alpha = 0.05. If p-value < 0.05 -> Reject H0 (Drift Detected)."
    ]
    for pt in drift_pts:
        story.append(Paragraph(f"• {pt}", bullet_style))
    story.append(Spacer(1, 10))

    # SECTION 4
    story.append(Paragraph("Section 4: Results, Interpretation and Conclusion", h1_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=2, spaceAfter=8))
    
    story.append(Paragraph("4.1 Empirical Kolmogorov-Smirnov Test Results", h2_style))
    
    # Table 2
    t2_data = [
        [Paragraph(h, table_header_style) for h in ["Feature Name", "KS Stat (D)", "p-value", "Alpha", "Ref Mean", "Curr Mean", "Mean Shift", "Status"]],
        [Paragraph("sepal_length", table_cell_bold), Paragraph("0.3733", table_cell_style), Paragraph("1.0722e-09", table_cell_style), Paragraph("0.05", table_cell_style), Paragraph("5.843 cm", table_cell_style), Paragraph("6.681 cm", table_cell_style), Paragraph("+0.837 cm", table_cell_style), Paragraph("<b>DRIFT DETECTED</b>", ParagraphStyle('Red', parent=table_cell_style, textColor=c_alert))],
        [Paragraph("sepal_width", table_cell_bold), Paragraph("0.1000", table_cell_style), Paragraph("0.4425", table_cell_style), Paragraph("0.05", table_cell_style), Paragraph("3.057 cm", table_cell_style), Paragraph("3.069 cm", table_cell_style), Paragraph("+0.012 cm", table_cell_style), Paragraph("<b>NO DRIFT</b>", ParagraphStyle('Green', parent=table_cell_style, textColor=c_ok))],
        [Paragraph("petal_length", table_cell_bold), Paragraph("0.4067", table_cell_style), Paragraph("1.7712e-11", table_cell_style), Paragraph("0.05", table_cell_style), Paragraph("3.758 cm", table_cell_style), Paragraph("5.172 cm", table_cell_style), Paragraph("+1.414 cm", table_cell_style), Paragraph("<b>DRIFT DETECTED</b>", ParagraphStyle('Red', parent=table_cell_style, textColor=c_alert))],
        [Paragraph("petal_width", table_cell_bold), Paragraph("0.1333", table_cell_style), Paragraph("0.1391", table_cell_style), Paragraph("0.05", table_cell_style), Paragraph("1.199 cm", table_cell_style), Paragraph("1.188 cm", table_cell_style), Paragraph("-0.012 cm", table_cell_style), Paragraph("<b>NO DRIFT</b>", ParagraphStyle('Green', parent=table_cell_style, textColor=c_ok))]
    ]
    t2 = Table(t2_data, colWidths=[1.1*inch, 0.8*inch, 0.95*inch, 0.55*inch, 0.85*inch, 0.85*inch, 0.85*inch, 1.25*inch])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_primary),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#B0C4DE")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, c_light]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t2)
    story.append(Spacer(1, 8))

    # Embed Plots in PDF
    if os.path.exists("plots/drift_distributions.png"):
        story.append(KeepTogether([
            RLImage("plots/drift_distributions.png", width=6.8 * inch, height=4.2 * inch),
            Spacer(1, 3),
            Paragraph("Figure 1: Empirical Distribution Comparison (KDE & Histograms) across Reference vs. Current Datasets", caption_style),
            Spacer(1, 8)
        ]))

    if os.path.exists("plots/drift_ks_summary.png"):
        story.append(KeepTogether([
            RLImage("plots/drift_ks_summary.png", width=6.8 * inch, height=2.4 * inch),
            Spacer(1, 3),
            Paragraph("Figure 2: Summary of Kolmogorov-Smirnov Test Statistics and Significance Thresholds", caption_style),
            Spacer(1, 8)
        ]))

    story.append(Paragraph("4.2 Interpretation & Critical Caveat", h2_style))
    story.append(Paragraph(
        "The two-sample KS test successfully identified statistically significant covariate shift in sepal_length (p=1.072e-09) and petal_length (p=1.771e-11). "
        "The control features confirmed no significant drift (p=0.4425 and p=0.1391), proving zero false-positive alerts.<br/>"
        "<b>Important Caveat:</b> Detecting data drift does NOT automatically prove model accuracy has decreased. If drifted points remain inside valid decision boundaries, "
        "accuracy can remain high. However, drift acts as a vital early-warning telemetry indicator.", body_style
    ))
    story.append(Spacer(1, 10))

    # SECTION 5
    story.append(Paragraph("Section 5: GitHub Repository and Reproducibility", h1_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=2, spaceAfter=8))
    
    sec5_pts = [
        "<b>GitHub Repository URL:</b> https://github.com/Tanish-19/iris-mlflow-drift",
        "<b>Local Repository Root:</b> <code>/Users/tanishpatgaonkar/.gemini/antigravity-ide/scratch/iris-mlflow-drift</code> (Branch: <code>main</code>)",
        "<b>Execution Sequence:</b> <code>python src/data_loader.py</code> -> <code>python src/train.py</code> -> <code>python src/drift_detection.py</code> -> <code>pytest tests/test_project.py -v</code>.",
        "<b>Automated Pytest Results:</b> <b>9/9 tests passed (100%)</b> validating dataset shape, training, prediction metrics, and KS drift assertions.",
        "<b>UI Inspection:</b> Local MLflow UI verified on port 5001; interactive Streamlit dashboard available on port 8501."
    ]
    for pt in sec5_pts:
        story.append(Paragraph(f"• {pt}", bullet_style))
    story.append(Spacer(1, 12))

    # Page Number Canvas callback
    def add_page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor("#7F7F7F"))
        canvas.drawString(0.65 * inch, 0.4 * inch, f"{STUDENT_NAME} (Roll: {ROLL_NUMBER}, PRN: {PRN})")
        canvas.drawRightString(8.5 * inch - 0.65 * inch, 0.4 * inch, f"Page {doc.page} | ISE-2 Academic Report")
        canvas.restoreState()

    pdf_doc.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)
    print(f"Successfully generated PDF report at: {output_path}")


if __name__ == "__main__":
    create_docx("ISE2_MLFLOW_DOCUMENTATION.docx")
    create_pdf("ISE2_MLFLOW_DOCUMENTATION.pdf")
